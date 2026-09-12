"""评估结果导出单测：双 sheet 结构、动态指标列、上下文格式化、错误判定口径。"""

import io

from yuxi.knowledge.eval.result_export import (
    build_run_results_workbook,
    collect_metric_columns,
    format_contexts_cell,
    is_error_item,
)


def _item(index: int, *, metrics: dict | None = None, chunks: list | None = None, gold=None) -> dict:
    return {
        "item_index": index,
        "query": f"问题{index}",
        "gold_chunk_ids": gold or [],
        "gold_answer": "标准答案文本" if index == 0 else None,
        "generated_answer": "生成答案文本",
        "retrieved_chunks": chunks
        if chunks is not None
        else [{"content": "上下文A", "metadata": {"chunk_id": "c1"}, "score": 0.9}],
        "metrics": metrics or {"recall@10": 1.0, "score": 1.0, "reasoning": "事实一致"},
    }


def test_format_contexts_cell_marks_gold_hits():
    item = {
        "gold_chunk_ids": ["c2"],
        "retrieved_chunks": [
            {"content": "第一条", "metadata": {"chunk_id": "c1"}, "score": 0.91},
            {"content": "第二条", "chunk_id": "c2", "score": 0.75},
        ],
    }

    text = format_contexts_cell(item)

    assert "【1】 c1" in text and "score=0.9100" in text
    assert "[命中标准块]" in text and "【2】 c2" in text
    assert "第一条" in text and "第二条" in text


def test_format_contexts_cell_empty():
    assert format_contexts_cell({"retrieved_chunks": []}) == ""
    assert format_contexts_cell({"retrieved_chunks": None}) == ""


def test_is_error_item_matches_frontend_thresholds():
    assert is_error_item({"metrics": {"score": 0.0}}) is True
    assert is_error_item({"metrics": {"recall@10": 0.2}}) is True
    assert is_error_item({"metrics": {"ragas_faithfulness": 0.1}}) is True
    assert is_error_item({"metrics": {"score": 1.0, "recall@10": 0.5}}) is False


def test_collect_metric_columns_excludes_judgement_and_sorts():
    items = [
        {"metrics": {"score": 1.0, "recall@10": 1.0, "ragas_faithfulness": 0.9}},
        {"metrics": {"reasoning": "x", "recall@1": 1.0, "f1@5": 0.5}},
    ]

    columns = collect_metric_columns(items)

    assert columns == ["recall@1", "recall@10", "f1@5", "ragas_faithfulness"]
    assert "score" not in columns and "reasoning" not in columns


def test_build_workbook_structure_and_content():
    run = {
        "run_id": "run_ab12cd34",
        "name": "eval-测试",
        "dataset_id": "dataset_x",
        "status": "completed",
        "started_at": "2026-09-08 10:00:00",
        "completed_at": "2026-09-08 10:05:00",
        "total_items": 2,
        "completed_items": 2,
        "overall_score": 0.85,
        "retrieval_config": {"eval_mode": "ragas"},
        "metrics": {
            "recall@10": 0.75,
            "ragas_faithfulness": 1.0,
            "metrics_meta": {
                "eval_mode": "ragas",
                "ragas": {"enabled": ["ragas_faithfulness"], "skipped": {}, "usage": {"llm_calls": 58}},
            },
        },
    }
    items = [
        _item(0, metrics={"recall@10": 1.0, "score": 1.0, "reasoning": "一致"}, gold=["c1"]),
        _item(1, metrics={"recall@10": 0.0, "score": 0.0, "reasoning": "矛盾"}),
    ]

    package = build_run_results_workbook(run=run, items=items)

    assert package["filename"] == "eval-results-eval-测试.xlsx"

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(package["content"]))
    assert workbook.sheetnames == ["运行汇总", "逐题明细"]

    summary = workbook["运行汇总"]
    summary_map = {row[0].value: row[1].value for row in summary.iter_rows(min_row=2, max_col=2)}
    assert summary_map["Run ID"] == "run_ab12cd34"
    assert summary_map["评估模式"] == "ragas"
    assert summary_map["指标 recall@10"] == 0.75
    assert "ragas_faithfulness" in str(summary_map["RAGAS 启用指标"])

    detail = workbook["逐题明细"]
    headers = [cell.value for cell in detail[1]]
    assert headers[:5] == [
        "序号",
        "问题 (question)",
        "标准答案 (ground truth)",
        "检索上下文 (contexts)",
        "生成答案 (answer)",
    ]
    assert "recall@10" in headers
    assert headers[-6:] == [
        "答案评判 (score)",
        "答案评判 (reasoning)",
        "命中标准块数",
        "是否错误",
        "标准块 ID (gold_chunk_ids)",
        "标签 (tags)",
    ]
    # 明细两行 + 表头
    assert detail.max_row == 3
    # 第 1 题命中 1 块、非错误；第 2 题未命中、score=0 判错误
    rows = list(detail.iter_rows(min_row=2, values_only=True))
    hit_col = headers.index("命中标准块数") + 1
    err_col = headers.index("是否错误") + 1
    assert rows[0][hit_col - 1] == 1 and rows[0][err_col - 1] == "否"
    assert rows[1][hit_col - 1] == 0 and rows[1][err_col - 1] == "是"


def test_build_workbook_without_judge_metrics():
    run = {"run_id": "run_x", "name": "n", "status": "completed", "metrics": {}, "retrieval_config": {}}
    items = [_item(0, metrics={"recall@10": 0.5})]

    package = build_run_results_workbook(run=run, items=items)

    from openpyxl import load_workbook

    detail = load_workbook(io.BytesIO(package["content"]))["逐题明细"]
    headers = [cell.value for cell in detail[1]]
    assert "答案评判 (score)" in headers  # 固定列保留，值为空
    score_col = headers.index("答案评判 (score)")
    assert detail.cell(row=2, column=score_col + 1).value in ("", None)
