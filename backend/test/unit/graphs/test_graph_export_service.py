"""图谱导出服务单测：往返自证（导出包直接喂回 v3 解析器）、identity 保持、xlsx 结构、
Neo4j 投影全量导出（流式 JSONL、事务内计数校验、规范层对账、单事务装配）。

核心断言：同一 kb_id 下，导出的 CSV 重新解析后，plan 中的 entity_id / triple_id /
evidence_id 与源数据完全一致——导入校验器就是导出格式的验收器。
"""

import csv
import hashlib
import io
import json
import zipfile

import pytest
from yuxi.knowledge.graphs.graph_export_service import (
    EvidenceTally,
    ProjectionTally,
    assemble_projection_package,
    build_evidence_workbook,
    build_projection_package,
    build_roundtrip_csvs,
    build_roundtrip_package,
    evidence_coverage,
    evidence_summary_header,
    export_projection_package,
    projection_chunk_row,
    projection_cypher,
    projection_decision_row,
    projection_evidence_row,
    projection_node_row,
    projection_relationship_row,
    reconcile_projection,
    reverse_map_node_type,
    verify_projection_counts,
    write_jsonl_member,
    write_jsonl_member_async,
)
from yuxi.knowledge.graphs.graph_utils import compute_entity_id, compute_triple_id, normalize_entity_name
from yuxi.knowledge.graphs.managed_import_parser import parse_managed_graph_import
from yuxi.utils import hashstr

KB = "kb_exporttest"

GENE_IDENTITY = "rap:os01g0100100"
PROCESS_IDENTITY = "name:endosperm development"
PHENOTYPE_IDENTITY = "name:chalky endosperm"
ALLELE_NAME = "os01g0100100 crispr"
ALLELE_IDENTITY = f"allele:{GENE_IDENTITY}:{normalize_entity_name(ALLELE_NAME)}"

GENE_ENTITY_ID = compute_entity_id(KB, GENE_IDENTITY, "Gene")
PROCESS_ENTITY_ID = compute_entity_id(KB, PROCESS_IDENTITY, "Process")
PHENOTYPE_ENTITY_ID = compute_entity_id(KB, PHENOTYPE_IDENTITY, "Phenotype")
ALLELE_ENTITY_ID = compute_entity_id(KB, ALLELE_IDENTITY, "AlleleMutant")

T_PROMOTES = compute_triple_id(KB, GENE_IDENTITY, "Gene", "PROMOTES_PROCESS", PROCESS_IDENTITY, "Process")
T_MUTANT = compute_triple_id(KB, ALLELE_IDENTITY, "AlleleMutant", "MUTANT_EFFECT", PHENOTYPE_IDENTITY, "Phenotype")
T_ALLELE_OF = compute_triple_id(KB, ALLELE_IDENTITY, "AlleleMutant", "ALLELE_OF", GENE_IDENTITY, "Gene")


def _expected_evidence_id(triple_id, literature_id, direction, directness, level, alignment, quote):
    identity = "|".join([KB, triple_id, literature_id or "", direction, directness, level, alignment, quote or ""])
    return hashstr(identity, length=32)


ALIGNED_QUOTE = "Os01g0100100 promotes endosperm development."
ROW_LEVEL_QUOTE = "q1\n\nq2"
EXPECTED_EV_ALIGNED = _expected_evidence_id(
    T_PROMOTES, "pmid:12345678", "POSITIVE", "DIRECT", "2", "ALIGNED", ALIGNED_QUOTE
)
EXPECTED_EV_ROW_LEVEL = _expected_evidence_id(T_MUTANT, None, "NEGATIVE", "DIRECT", "2", "ROW_LEVEL", ROW_LEVEL_QUOTE)


def _entity(entity_id, identity, label, name, *, attributes=None):
    return {
        "entity_id": entity_id,
        "kb_id": KB,
        "canonical_identity": identity,
        "normalized_name": normalize_entity_name(name),
        "label": label,
        "name": name,
        "attributes": attributes or {},
        "content": name,
    }


def _triple(triple_id, source_entity_id, target_entity_id, relation_type, **kwargs):
    return {
        "triple_id": triple_id,
        "kb_id": KB,
        "source_entity_id": source_entity_id,
        "target_entity_id": target_entity_id,
        "relation_type": relation_type,
        "content": f"{relation_type}",
        "support_count": kwargs.get("support_count", 1),
        "literature_count": kwargs.get("literature_count", 1),
        "best_evidence_level": kwargs.get("best_evidence_level", "2"),
        "consensus_direction": kwargs.get("consensus_direction", "POSITIVE"),
    }


def _evidence(evidence_id, triple_id, **kwargs):
    return {
        "evidence_id": evidence_id,
        "triple_id": triple_id,
        "literature_id": kwargs.get("literature_id"),
        "pmid": kwargs.get("pmid"),
        "doi": kwargs.get("doi"),
        "identifier_status": kwargs.get("identifier_status", "VALID"),
        "direction": kwargs.get("direction", "POSITIVE"),
        "directness": kwargs.get("directness", "DIRECT"),
        "assertion_status": "ASSERTED",
        "evidence_level": kwargs.get("evidence_level", "2"),
        "evidence_quote": kwargs.get("evidence_quote"),
        "evidence_alignment_status": kwargs.get("evidence_alignment_status", "ALIGNED"),
        "outcome_class": kwargs.get("outcome_class", "OTHER"),
        "metadata_json": kwargs.get("metadata_json", {}),
    }


@pytest.fixture()
def sample_snapshot():
    entities = [
        _entity(
            GENE_ENTITY_ID,
            GENE_IDENTITY,
            "Gene",
            "Os01g0100100",
            attributes={
                "source": "managed_csv_import",
                "external_ids": ["Os01g0100100"],
                "aliases": ["Os01g0100100"],
                "rap_ids": ["Os01g0100100"],
                "msu_ids": ["LOC_Os01g01009.1"],
                "canonical_identity": GENE_IDENTITY,
                "gene_status": "confirmed",
            },
        ),
        _entity(PROCESS_ENTITY_ID, PROCESS_IDENTITY, "Process", "endosperm development"),
        _entity(PHENOTYPE_ENTITY_ID, PHENOTYPE_IDENTITY, "Phenotype", "chalky endosperm"),
        _entity(
            ALLELE_ENTITY_ID,
            ALLELE_IDENTITY,
            "AlleleMutant",
            ALLELE_NAME,
            attributes={
                "external_ids": [ALLELE_NAME],
                "aliases": [ALLELE_NAME],
                "rap_ids": [],
                "msu_ids": [],
                "canonical_identity": ALLELE_IDENTITY,
                "parent_gene_identity": GENE_IDENTITY,
            },
        ),
    ]
    triples = [
        _triple(T_PROMOTES, GENE_ENTITY_ID, PROCESS_ENTITY_ID, "PROMOTES_PROCESS"),
        _triple(T_MUTANT, ALLELE_ENTITY_ID, PHENOTYPE_ENTITY_ID, "MUTANT_EFFECT"),
        # 派生关系：导出应跳过，重导入由 normalizer 重建
        _triple(
            T_ALLELE_OF,
            ALLELE_ENTITY_ID,
            GENE_ENTITY_ID,
            "ALLELE_OF",
            support_count=0,
            literature_count=0,
            best_evidence_level=None,
            consensus_direction="UNKNOWN",
        ),
    ]
    evidence = [
        _evidence(
            "ev_aligned_1", T_PROMOTES, pmid="12345678", literature_id="pmid:12345678", evidence_quote=ALIGNED_QUOTE
        ),
        _evidence(
            "ev_rowlevel_1",
            T_MUTANT,
            evidence_alignment_status="ROW_LEVEL",
            direction="NEGATIVE",
            evidence_quote=ROW_LEVEL_QUOTE,
            metadata_json={"pmids": ["11111111", "22222222"], "dois": ["10.1234/test.001"]},
        ),
    ]
    aliases = {GENE_ENTITY_ID: ["OS01G0100100"]}
    return {"entities": entities, "triples": triples, "evidence": evidence, "aliases": aliases}


def _read_csv_rows(data: bytes) -> list[dict[str, str]]:
    reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig"), newline=""))
    return list(reader)


# ---------------------------------------------------------------- 反向映射 ----


def test_reverse_map_node_type_gene_three_states():
    assert reverse_map_node_type("Gene", "confirmed") == "RICE_GENE"
    assert reverse_map_node_type("Gene", "candidate") == "RICE_GENE_CANDIDATE"
    assert reverse_map_node_type("Gene", "unspecified") == "GENE"
    assert reverse_map_node_type("Gene", None) == "GENE"
    assert reverse_map_node_type("AlleleMutant", None) == "ALLELE_MUTANT"
    assert reverse_map_node_type("QTL", None) == "QTL_LOCUS"
    assert reverse_map_node_type("Compound", None) is None


# ------------------------------------------------------------ 往返自证 ----


def test_roundtrip_reparse_preserves_identity(sample_snapshot):
    package = build_roundtrip_package(
        kb_id=KB,
        entities=sample_snapshot["entities"],
        triples=sample_snapshot["triples"],
        evidence=sample_snapshot["evidence"],
        aliases_by_entity=sample_snapshot["aliases"],
        exported_by="tester",
    )

    with zipfile.ZipFile(io.BytesIO(package["content"])) as archive:
        names = set(archive.namelist())
        nodes_csv = archive.read("nodes.csv")
        relationships_csv = archive.read("relationships.csv")
        manifest = json.loads(archive.read("manifest.json").decode("utf-8"))

    assert names == {"nodes.csv", "relationships.csv", "manifest.json"}
    assert manifest["schema_version"] == "rice-endosperm-csv-v3"
    assert manifest["reparse_self_check"]["status"] == "READY"
    assert manifest["reparse_self_check"]["identity_preserved"] is True

    reparsed = parse_managed_graph_import(kb_id=KB, nodes_bytes=nodes_csv, relationships_bytes=relationships_csv)
    assert reparsed["valid"] is True
    assert reparsed["blockers"] == []

    plan = reparsed["plan"]
    # identity 保持：同 kb 重导入 entity_id/triple_id 完全一致（含派生 ALLELE_OF 重建）
    assert {entity["entity_id"] for entity in plan["entities"]} == {
        GENE_ENTITY_ID,
        PROCESS_ENTITY_ID,
        PHENOTYPE_ENTITY_ID,
        ALLELE_ENTITY_ID,
    }
    assert {triple["triple_id"] for triple in plan["triples"]} == {T_PROMOTES, T_MUTANT, T_ALLELE_OF}
    # 证据逐条还原：ALIGNED 单条 + ROW_LEVEL 捆绑包重导入后 evidence_id 与解析器算法一致
    assert {item["evidence_id"] for item in plan["evidence"]} == {EXPECTED_EV_ALIGNED, EXPECTED_EV_ROW_LEVEL}


def test_roundtrip_node_rows(sample_snapshot):
    built = build_roundtrip_csvs(
        entities=sample_snapshot["entities"],
        triples=sample_snapshot["triples"],
        evidence=sample_snapshot["evidence"],
        aliases_by_entity=sample_snapshot["aliases"],
    )
    node_rows = _read_csv_rows(built["nodes_csv"])
    rice_gene_rows = [row for row in node_rows if row["node_type"] == "RICE_GENE"]

    assert any(row["rap_id"] == "Os01g0100100" for row in rice_gene_rows)
    # 多 registry id 拆行，msu 单独一行，node_id 相同
    gene_msu_rows = [
        row for row in node_rows if row["msu_id"] == "LOC_Os01g01009.1" and row["node_id"] == "Os01g0100100"
    ]
    assert len(gene_msu_rows) == 1
    # allele 行挂靠父基因 node_id（union-find 并组 -> 语义拆分重建 parent_identity）
    allele_rows = [row for row in node_rows if row["node_type"] == "ALLELE_MUTANT"]
    assert len(allele_rows) == 1
    assert allele_rows[0]["node_id"] == "Os01g0100100"
    assert allele_rows[0]["name"] == ALLELE_NAME
    # 别名附加列（导入器忽略）
    assert any(row["aliases"] == "OS01G0100100" for row in rice_gene_rows)

    relation_rows = _read_csv_rows(built["relationships_csv"])
    # ALLELE_OF 派生关系不导出
    assert all(row["relation_type"] != "ALLELE_OF" for row in relation_rows)
    aligned = next(row for row in relation_rows if row["pmids"] == "12345678")
    assert aligned["relation_type"] == "PROMOTES_PROCESS"
    row_level = next(row for row in relation_rows if row["pmids"] == "11111111|22222222")
    assert row_level["dois"] == "10.1234/test.001" and row_level["evidence_quotes"] == "q1\n\nq2"


def test_roundtrip_skips_unknown_labels_and_dangling_relations(sample_snapshot):
    unknown = _entity("e_compound", "name:starch", "Compound", "starch")
    sample_snapshot["entities"].append(unknown)
    sample_snapshot["triples"].append(_triple("t_dangling_1", GENE_ENTITY_ID, "e_compound", "REQUIRED_FOR"))

    package = build_roundtrip_package(
        kb_id=KB,
        entities=sample_snapshot["entities"],
        triples=sample_snapshot["triples"],
        evidence=sample_snapshot["evidence"],
        aliases_by_entity=sample_snapshot["aliases"],
    )
    manifest = package["manifest"]
    assert manifest["skipped"]["unsupported_label_entities"] == [
        {"entity_id": "e_compound", "label": "Compound", "name": "starch"}
    ]
    assert manifest["skipped"]["triples_with_skipped_endpoint"] == 1


def test_roundtrip_empty_data_produces_header_only_csvs():
    built = build_roundtrip_csvs(entities=[], triples=[], evidence=[])
    assert _read_csv_rows(built["nodes_csv"]) == []
    assert _read_csv_rows(built["relationships_csv"]) == []


# ------------------------------------------------------------- xlsx ----


def test_evidence_workbook_structure(sample_snapshot):
    package = build_evidence_workbook(
        kb_id=KB,
        entities=sample_snapshot["entities"],
        triples=sample_snapshot["triples"],
        evidence=sample_snapshot["evidence"],
        aliases_by_entity=sample_snapshot["aliases"],
    )
    assert package["filename"].endswith(".xlsx")

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(package["content"]))
    assert workbook.sheetnames == ["实体", "三元组", "证据明细"]
    entity_sheet = workbook["实体"]
    assert entity_sheet.max_row == len(sample_snapshot["entities"]) + 1
    assert entity_sheet["A1"].value == "entity_id"
    triple_sheet = workbook["三元组"]
    assert triple_sheet.max_row == len(sample_snapshot["triples"]) + 1
    evidence_sheet = workbook["证据明细"]
    assert evidence_sheet.max_row == len(sample_snapshot["evidence"]) + 1
    headers = [cell.value for cell in evidence_sheet[1]]
    assert "experimental_subject_type" in headers and "claim_eligible" in headers
    # source 列渲染为 名称(标签)
    assert triple_sheet.cell(row=2, column=2).value == "Os01g0100100(Gene)"


# ------------------------------------------------------------ projection ----

PROJ_KB = "kb_projtest"
PROJ_CYPHER = projection_cypher(PROJ_KB)


def _projection_records():
    """两轨混合投影：托管实体 + chunk 抽取实体、Chunk 节点、同 triple 的两条 RELATION 边、一条 MENTIONS。"""
    gene = {
        "element_id": "4:store:1",
        "labels": ["Entity", "MilvusKB", PROJ_KB],
        "properties": {
            "entity_id": "ent_gene",
            "kb_id": PROJ_KB,
            "name": "OsNF-YB1",
            "label": "Gene",
            "attributes": "{}",
            "managed_projection": True,
        },
    }
    phenotype = {
        "element_id": "4:store:2",
        "labels": ["MilvusKB", "Entity", PROJ_KB],
        "properties": {"entity_id": "ent_phen", "kb_id": PROJ_KB, "name": "chalky endosperm", "label": "Phenotype"},
    }
    chunk = {
        "element_id": "4:store:3",
        "labels": ["Chunk", "MilvusKB", PROJ_KB],
        "properties": {
            "chunk_id": "chunk_1",
            "kb_id": PROJ_KB,
            "file_id": "file_1",
            "content_preview": 'line1\n"quoted", 逗号',
        },
    }

    def edge(element_id, rel_type, properties, start, end):
        return {
            "element_id": element_id,
            "relationship_type": rel_type,
            "properties": properties,
            "start_element_id": start["element_id"],
            "start_labels": start["labels"],
            "start_properties": start["properties"],
            "end_element_id": end["element_id"],
            "end_labels": end["labels"],
            "end_properties": end["properties"],
        }

    relationships = [
        edge(
            "5:store:10",
            "RELATION",
            {"triple_id": "tri_1", "type": "PROMOTES_PHENOTYPE", "chunk_id": "c1"},
            gene,
            phenotype,
        ),
        edge(
            "5:store:11",
            "RELATION",
            {"triple_id": "tri_1", "type": "PROMOTES_PHENOTYPE", "chunk_id": "c2"},
            gene,
            phenotype,
        ),
        edge("5:store:12", "MENTIONS", {"chunk_id": "chunk_1", "file_id": "file_1"}, chunk, gene),
    ]
    return [gene, phenotype, chunk], relationships


def _fake_run(nodes, relationships, *, node_counts=None, relationship_counts=None):
    responses = {
        PROJ_CYPHER["node_counts"]: [node_counts or {"total": 3, "entity": 2, "chunk": 1}],
        PROJ_CYPHER["relationship_counts"]: relationship_counts
        or [{"relationship_type": "RELATION", "total": 2}, {"relationship_type": "MENTIONS", "total": 1}],
        PROJ_CYPHER["nodes"]: nodes,
        PROJ_CYPHER["relationships"]: relationships,
    }
    return lambda query: iter(responses[query])


def _matching_reference():
    return {
        "entity_ids": {"ent_gene", "ent_phen"},
        "triple_ids": {"tri_1"},
        "mention_keys": {"chunk_1->ent_gene"},
        "chunk_ids": {"chunk_1", "chunk_pending"},
        "graph_indexed_chunk_ids": {"chunk_1"},
        "kb_name": "投影测试库",
    }


def _read_jsonl(archive: zipfile.ZipFile, name: str) -> list[dict]:
    return [json.loads(line) for line in archive.read(name).decode("utf-8").splitlines() if line]


def test_projection_node_row_classifies_and_extracts_business_id():
    nodes, _ = _projection_records()
    gene_row = projection_node_row(nodes[0])
    chunk_row = projection_node_row(nodes[2])
    other_row = projection_node_row({"element_id": "x", "labels": ["MilvusKB", PROJ_KB], "properties": {"name": "?"}})

    assert gene_row["node_class"] == "entity" and gene_row["business_id"] == "ent_gene"
    assert gene_row["labels"] == sorted(["Entity", "MilvusKB", PROJ_KB])
    assert gene_row["properties"]["managed_projection"] is True
    assert chunk_row["node_class"] == "chunk" and chunk_row["business_id"] == "chunk_1"
    assert other_row["node_class"] == "other" and other_row["business_id"] is None


def test_projection_relationship_row_business_ids():
    _, relationships = _projection_records()
    relation_row = projection_relationship_row(relationships[0])
    mention_row = projection_relationship_row(relationships[2])
    orphan = projection_relationship_row({**relationships[0], "properties": {"type": "X"}})

    assert relation_row["business_id"] == "tri_1"
    assert relation_row["start"] == {"element_id": "4:store:1", "node_class": "entity", "business_id": "ent_gene"}
    assert relation_row["end"]["business_id"] == "ent_phen"
    assert mention_row["business_id"] == "chunk_1->ent_gene"
    assert mention_row["start"]["node_class"] == "chunk" and mention_row["end"]["node_class"] == "entity"
    assert orphan["business_id"] is None


def test_write_jsonl_member_streams_rows_and_hashes_member_bytes():
    nodes, _ = _projection_records()
    buffer = io.BytesIO()
    seen = []
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        stat = write_jsonl_member(archive, "nodes.jsonl", iter(nodes), projection_node_row, seen.append)
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        raw = archive.read("nodes.jsonl")
        rows = _read_jsonl(archive, "nodes.jsonl")

    assert stat["rows"] == 3 and len(seen) == 3
    assert stat["bytes"] == len(raw)
    assert stat["sha256"] == hashlib.sha256(raw).hexdigest()
    # 换行 / 引号 / 逗号 在 JSONL 中无损往返
    assert rows[2]["properties"]["content_preview"] == 'line1\n"quoted", 逗号'


def test_projection_package_exports_every_node_and_relationship():
    nodes, relationships = _projection_records()
    package = assemble_projection_package(
        kb_id=PROJ_KB,
        run=_fake_run(nodes, relationships),
        reference=_matching_reference(),
        exported_by="tester",
    )

    assert package["filename"] == f"graph-projection-{PROJ_KB}.zip"
    assert package["media_type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(package["content"])) as archive:
        assert set(archive.namelist()) == {"nodes.jsonl", "relationships.jsonl", "manifest.json"}
        node_rows = _read_jsonl(archive, "nodes.jsonl")
        relationship_rows = _read_jsonl(archive, "relationships.jsonl")
        manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
        assert manifest["files"]["nodes.jsonl"]["sha256"] == hashlib.sha256(archive.read("nodes.jsonl")).hexdigest()

    # 四类元素一个不落
    assert {row["business_id"] for row in node_rows if row["node_class"] == "entity"} == {"ent_gene", "ent_phen"}
    assert {row["business_id"] for row in node_rows if row["node_class"] == "chunk"} == {"chunk_1"}
    assert [row["relationship_type"] for row in relationship_rows] == ["RELATION", "RELATION", "MENTIONS"]

    assert manifest["schema_version"] == "neo4j-projection-jsonl-v1"
    assert manifest["kb_name"] == "投影测试库" and manifest["exported_by"] == "tester"
    assert manifest["source"]["role"] == "rebuildable_projection"
    assert manifest["counts"]["nodes"] == {"total": 3, "entity": 2, "chunk": 1, "other": 0}
    assert manifest["counts"]["relationships"] == {"total": 3, "MENTIONS": 1, "RELATION": 2}
    # 同 triple 两条边去重后只算一个三元组
    assert manifest["counts"]["distinct_business_ids"] == {"entities": 2, "chunks": 1, "triples": 1, "mentions": 1}
    assert manifest["verification"]["status"] == "PASSED"
    assert manifest["reconciliation"]["matched"] is True
    assert manifest["files"]["relationships.jsonl"]["rows"] == 3
    assert manifest["scope"]["relationship_types"] == ["MENTIONS", "RELATION"]


def test_projection_reconciliation_reports_drift_in_manifest():
    nodes, relationships = _projection_records()
    reference = _matching_reference()
    reference["entity_ids"] = {"ent_gene", "ent_missing"}  # 规范层多一个未投影实体
    reference["mention_keys"] = set()  # 投影里多出一条提及
    package = assemble_projection_package(kb_id=PROJ_KB, run=_fake_run(nodes, relationships), reference=reference)
    reconciliation = package["manifest"]["reconciliation"]

    assert reconciliation["matched"] is False
    assert reconciliation["entities"]["missing_in_projection"] == ["ent_missing"]
    assert reconciliation["entities"]["extra_in_projection"] == ["ent_phen"]
    assert reconciliation["mentions"]["extra_in_projection"] == ["chunk_1->ent_gene"]
    assert reconciliation["triples"]["missing_in_projection_count"] == 0
    # 未索引的 pending chunk 不算缺失，已索引的 chunk_1 有节点 -> chunk 维度本身干净
    assert reconciliation["chunks"]["missing_in_projection_count"] == 0
    assert reconciliation["chunks"]["extra_in_projection_count"] == 0
    assert any("对账未通过" in note for note in package["manifest"]["notes"])


def test_projection_chunk_reconciliation_distinguishes_stale_missing_and_unflagged():
    tally = ProjectionTally()
    for chunk_id in ("chunk_stale", "chunk_ok", "chunk_unflagged"):
        tally.add_node({"node_class": "chunk", "business_id": chunk_id})
    reference = {
        # PG 里 chunk_stale 已不存在（文件重新解析后级联删除）；chunk_lost 已索引却没有节点
        "chunk_ids": {"chunk_ok", "chunk_unflagged", "chunk_lost", "chunk_pending"},
        "graph_indexed_chunk_ids": {"chunk_ok", "chunk_lost"},
    }

    chunks = reconcile_projection(reference, tally)["chunks"]

    assert chunks["canonical_count"] == 2 and chunks["canonical_total_chunks"] == 4
    assert chunks["projected_count"] == 3
    assert chunks["missing_in_projection"] == ["chunk_lost"]
    assert chunks["extra_in_projection"] == ["chunk_stale"]
    assert chunks["projected_but_not_flagged"] == ["chunk_unflagged"]
    assert reconcile_projection(reference, tally)["matched"] is False


def test_projection_count_mismatch_rejects_package():
    nodes, relationships = _projection_records()
    run = _fake_run(nodes, relationships, node_counts={"total": 4, "entity": 3, "chunk": 1})
    with pytest.raises(RuntimeError, match="计数校验失败"):
        assemble_projection_package(kb_id=PROJ_KB, run=run, reference=_matching_reference())

    run = _fake_run(nodes, relationships, relationship_counts=[{"relationship_type": "RELATION", "total": 2}])
    with pytest.raises(RuntimeError, match="MENTIONS 关系 期望 0 实际 1"):
        assemble_projection_package(kb_id=PROJ_KB, run=run, reference=_matching_reference())


def test_projection_empty_store_raises_not_found_message():
    run = _fake_run([], [], node_counts={"total": 0, "entity": 0, "chunk": 0}, relationship_counts=[])
    with pytest.raises(ValueError, match="没有可导出"):
        assemble_projection_package(kb_id=PROJ_KB, run=run, reference=_matching_reference())


def test_verify_projection_counts_passes_when_tally_matches():
    tally = ProjectionTally()
    tally.add_node({"node_class": "entity", "business_id": "e1"})
    tally.add_node({"node_class": "other", "business_id": None})
    tally.add_relationship({"relationship_type": "RELATION", "business_id": None})
    result = verify_projection_counts({"total": 2, "entity": 1, "chunk": 0}, {"RELATION": 1}, tally)

    assert result["status"] == "PASSED"
    assert tally.nodes_without_business_id == 1 and tally.relationships_without_business_id == 1
    assert reconcile_projection({"entity_ids": {"e1"}}, tally)["matched"] is True


class _FakeRecord:
    def __init__(self, data):
        self._data = data

    def data(self):
        return self._data


class _FakeTransaction:
    def __init__(self, run):
        self._run = run
        self.queries = []
        self.committed = False
        self.exited = False

    def run(self, query):
        self.queries.append(query)
        return [_FakeRecord(item) for item in self._run(query)]

    def commit(self):
        self.committed = True

    def __enter__(self):
        return self

    def __exit__(self, *_exc):
        self.exited = True
        return False


class _FakeSession:
    def __init__(self, transaction):
        self._transaction = transaction
        self.transactions_begun = 0

    def begin_transaction(self):
        self.transactions_begun += 1
        return self._transaction

    def __enter__(self):
        return self

    def __exit__(self, *_exc):
        return False


class _FakeDriver:
    def __init__(self, transaction):
        self.session_instance = _FakeSession(transaction)

    def session(self):
        return self.session_instance


def test_build_projection_package_runs_all_queries_in_one_committed_transaction():
    nodes, relationships = _projection_records()
    transaction = _FakeTransaction(_fake_run(nodes, relationships))
    driver = _FakeDriver(transaction)

    package = build_projection_package(kb_id=PROJ_KB, driver=driver, reference=_matching_reference())

    assert driver.session_instance.transactions_begun == 1
    assert transaction.queries == [
        PROJ_CYPHER["node_counts"],
        PROJ_CYPHER["relationship_counts"],
        PROJ_CYPHER["nodes"],
        PROJ_CYPHER["relationships"],
    ]
    assert transaction.committed and transaction.exited
    assert package["manifest"]["counts"]["relationships"]["total"] == 3


def test_build_projection_package_rejects_unsafe_kb_label():
    with pytest.raises(ValueError, match="非法 Neo4j 标签"):
        assemble_projection_package(kb_id="kb-bad`) DETACH DELETE", run=lambda _q: iter([]), reference={})


# ------------------------------------------- projection v2：原文证据成员 ----

CHUNK_CONTENT = (
    "【章节】3.2 胚乳发育调控\n"
    "【文献】OsNF-YB1 调控胚乳发育\n"
    "【页码】4\n"
    "OsNF-YB1 通过结合 OsDOG1L 启动子促进胚乳细胞增殖。"
)
CHUNK_QUOTE = "OsNF-YB1 通过结合 OsDOG1L 启动子促进胚乳细胞增殖。"


def _mention_record(**overrides):
    """PG 统一证据行（knowledge_graph_repository._evidence_row 的形状）。"""
    record = {
        "kind": "triple",
        "source": "chunk_mention",
        "target_id": "tri_1",
        "edge_business_id": "tri_1",
        "chunk_id": "c1",
        "file_id": "file_1",
        "filename": "rice.pdf",
        "chunk_index": 0,
        "quote": CHUNK_QUOTE,
        "quote_start_char": None,
        "extractor_type": "llm_scientific",
        "confidence": 0.87,
        "hedge": False,
        "context": {"condition": "OE"},
        "trigger_verified": True,
        "trigger_term": "促进",
        "verifier_confirmed": True,
        "pinned_by": None,
        "pinned_at": None,
        "review_status": "APPROVED",
        "pmid": None,
        "doi": None,
        "chunk_content": CHUNK_CONTENT,
        "source_provenance": None,
    }
    record.update(overrides)
    return record


def _chunk_record(chunk_id: str, **overrides):
    record = {
        "chunk_id": chunk_id,
        "file_id": "file_1",
        "filename": "rice.pdf",
        "chunk_index": 0,
        "content": CHUNK_CONTENT,
        "start_char_pos": 0,
        "end_char_pos": len(CHUNK_CONTENT),
        "source_provenance": {"page_numbers": [4]},
    }
    record.update(overrides)
    return record


def test_projection_evidence_row_verifies_quote_and_parses_provenance():
    row = projection_evidence_row(_mention_record())

    assert row["verification"] == "OK"
    assert row["quote"] == CHUNK_QUOTE
    assert row["quote_start_char"] == CHUNK_CONTENT.index(CHUNK_QUOTE)
    # 【章节】【文献】【页码】前缀标记行解析为离线可读来源
    assert row["section"] == "3.2 胚乳发育调控"
    assert row["literature"] == "OsNF-YB1 调控胚乳发育"
    assert row["page"] == "4"
    assert (row["kind"], row["business_id"], row["edge_business_id"]) == ("triple", "tri_1", "tri_1")
    assert row["review_status"] == "APPROVED" and row["filename"] == "rice.pdf"
    assert row["confidence"] == 0.87 and row["trigger_verified"] is True


def test_projection_evidence_row_marks_degraded_missing_and_stored_offset():
    degraded = projection_evidence_row(_mention_record(quote="这句话不在原文里。"))
    collapsed = projection_evidence_row(_mention_record(quote=CHUNK_QUOTE.replace(" ", "\n")))
    legacy = projection_evidence_row(_mention_record(quote=None))
    stored = projection_evidence_row(_mention_record(quote_start_char=7))

    assert degraded["verification"] == "DEGRADED"
    # 仅空白差异仍算 OK，但不给偏移（否则高亮会对错位置）
    assert collapsed["verification"] == "OK" and collapsed["quote_start_char"] is None
    assert legacy["verification"] == "MISSING" and legacy["quote"] is None
    # 库里已存的偏移优先于导出时重算（与面板同源）
    assert stored["quote_start_char"] == 7


def test_projection_evidence_row_marks_managed_import_unverifiable():
    row = projection_evidence_row(
        {
            "kind": "triple",
            "source": "relation_evidence",
            "target_id": "tri_1",
            "edge_business_id": "tri_1",
            "quote": "Os01g0100100 promotes endosperm development.",
            "pmid": "12345678",
            "review_status": "CANONICAL",
            "chunk_content": None,
        }
    )
    empty = projection_evidence_row({"kind": "triple", "source": "relation_evidence", "quote": None})

    assert row["verification"] == "UNVERIFIABLE" and row["chunk_id"] is None
    assert row["pmid"] == "12345678" and row["filename"] is None
    assert empty["verification"] == "MISSING"


def test_projection_chunk_row_hashes_text_and_can_drop_it():
    record = _chunk_record("c1", start_char_pos=100, end_char_pos=200)
    with_text = projection_chunk_row(record, include_text=True)
    without_text = projection_chunk_row(record, include_text=False)

    assert with_text["content"] == CHUNK_CONTENT
    assert with_text["content_sha256"] == hashlib.sha256(CHUNK_CONTENT.encode("utf-8")).hexdigest()
    assert with_text["content_length"] == len(CHUNK_CONTENT)
    assert with_text["start_char_pos"] == 100 and with_text["end_char_pos"] == 200
    # 大库逃生口：丢正文但保留哈希，离线仍可校验「引文与所声称的段落一致」
    assert without_text["content"] is None
    assert without_text["content_sha256"] == with_text["content_sha256"]


def test_projection_decision_row_keeps_audit_fields():
    row = projection_decision_row(
        {
            "decision_id": "dec_1",
            "target_kind": "triple",
            "target_id": "tri_9",
            "action": "REJECT",
            "reason": "引文与原文不符",
            "actor_uid": "u_admin",
            "version": 3,
            "pinned_chunk_id": "c1",
            "pinned_quote": CHUNK_QUOTE,
            "created_at": "2026-01-01T00:00:00+00:00",
            "updated_at": "2026-01-02T00:00:00+00:00",
        }
    )

    assert row["action"] == "REJECT" and row["reason"] == "引文与原文不符"
    assert row["actor_uid"] == "u_admin" and row["pinned_quote"] == CHUNK_QUOTE
    assert row["version"] == 3 and row["updated_at"] == "2026-01-02T00:00:00+00:00"


def _coverage_tally():
    """与 _projection_records() 同形的投影侧旁路统计：两条 RELATION 边 + 一条 MENTIONS 边。"""
    tally = ProjectionTally()
    tally.add_relationship({"relationship_type": "RELATION", "business_id": "tri_1", "properties": {"chunk_id": "c1"}})
    tally.add_relationship({"relationship_type": "RELATION", "business_id": "tri_1", "properties": {"chunk_id": "c2"}})
    tally.add_relationship({"relationship_type": "MENTIONS", "business_id": "chunk_1->ent_gene"})
    return tally


def _collect_evidence(records):
    tally = EvidenceTally()
    for record in records:
        tally.add(projection_evidence_row(record))
    return tally


def _projection_evidence_bundle():
    """与 _projection_records() 边一一对应的 PG 证据包（含段落全文与一条人工决策）。"""
    return {
        "rows": [
            _mention_record(),
            _mention_record(chunk_id="c2"),
            _mention_record(
                kind="entity", target_id="ent_gene", chunk_id="chunk_1", edge_business_id="chunk_1->ent_gene"
            ),
        ],
        "chunks": [_chunk_record("c1"), _chunk_record("c2"), _chunk_record("chunk_1")],
        "decisions": [
            {
                "target_kind": "triple",
                "target_id": "tri_9",
                "action": "REJECT",
                "reason": "引文与原文不符",
                "actor_uid": "u_admin",
                "version": 2,
                "pinned_chunk_id": "c1",
                "pinned_quote": CHUNK_QUOTE,
                "updated_at": "2026-02-01T00:00:00+00:00",
            }
        ],
    }


def test_projection_package_with_evidence_members_is_v2_and_keeps_v1_rows():
    nodes, relationships = _projection_records()
    reference = _matching_reference()
    light = assemble_projection_package(kb_id=PROJ_KB, run=_fake_run(nodes, relationships), reference=reference)
    full = assemble_projection_package(
        kb_id=PROJ_KB,
        run=_fake_run(nodes, relationships),
        reference=reference,
        exported_by="tester",
        evidence_bundle=_projection_evidence_bundle(),
    )

    with zipfile.ZipFile(io.BytesIO(light["content"])) as archive:
        assert set(archive.namelist()) == {"nodes.jsonl", "relationships.jsonl", "manifest.json"}
        light_nodes = archive.read("nodes.jsonl")
        light_relationships = archive.read("relationships.jsonl")

    with zipfile.ZipFile(io.BytesIO(full["content"])) as archive:
        assert set(archive.namelist()) == {
            "nodes.jsonl",
            "relationships.jsonl",
            "evidence.jsonl",
            "chunks.jsonl",
            "decisions.jsonl",
            "manifest.json",
        }
        # 加法演进的硬保证：v1 两个成员逐字节不变，旧消费者不受影响
        assert archive.read("nodes.jsonl") == light_nodes
        assert archive.read("relationships.jsonl") == light_relationships
        evidence_rows = _read_jsonl(archive, "evidence.jsonl")
        chunk_rows = _read_jsonl(archive, "chunks.jsonl")
        decision_rows = _read_jsonl(archive, "decisions.jsonl")
        manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
        # 包内自证：manifest 的 sha256 必须等于解压后成员内容的哈希
        for name in ("nodes.jsonl", "relationships.jsonl", "evidence.jsonl", "chunks.jsonl", "decisions.jsonl"):
            assert manifest["files"][name]["sha256"] == hashlib.sha256(archive.read(name)).hexdigest()

    assert manifest["schema_version"] == "neo4j-projection-jsonl-v2"
    assert manifest["scope"]["consistency"] == "neo4j_single_read_transaction + postgresql_read_after"
    evidence = manifest["evidence"]
    assert evidence["counts"] == {
        "evidence_rows": 3,
        "by_kind": {"triple": 2, "entity": 1},
        "by_source": {"chunk_mention": 3},
        "referenced_chunks": 3,
        "chunk_rows": 3,
        "missing_chunk_rows": 0,
        "decision_rows": 1,
    }
    assert evidence["verification"] == {"OK": 3}
    assert evidence["review_status"] == {"APPROVED": 3}
    assert evidence["coverage"]["evidence_complete"] is True
    assert evidence["include_chunk_text"] is True
    assert len(evidence_rows) == 3 and len(chunk_rows) == 3
    # 关联键可精确 join 回关系行：entity mention 的 edge_business_id 就是 MENTIONS 边键
    assert {row["edge_business_id"] for row in evidence_rows} == {"tri_1", "chunk_1->ent_gene"}
    assert decision_rows[0]["action"] == "REJECT" and decision_rows[0]["pinned_quote"] == CHUNK_QUOTE
    assert chunk_rows[0]["content"] == CHUNK_CONTENT
    # 旧的那条「无法还原逐条证据」声明已被如实描述替换
    assert all("无法还原" not in note for note in manifest["notes"])
    assert any("chunks.jsonl 提供段落全文与哈希" in note for note in manifest["notes"])
    assert any("已拒绝" in note for note in manifest["notes"])


def test_projection_lightweight_package_stays_v1():
    nodes, relationships = _projection_records()
    package = assemble_projection_package(
        kb_id=PROJ_KB, run=_fake_run(nodes, relationships), reference=_matching_reference()
    )
    manifest = package["manifest"]

    assert manifest["schema_version"] == "neo4j-projection-jsonl-v1"
    assert "evidence" not in manifest
    assert manifest["scope"]["consistency"] == "single_read_transaction"
    assert evidence_summary_header(package) is None
    assert any("仅结构轻量导出" in note for note in manifest["notes"])


def test_projection_evidence_chunk_text_can_be_dropped():
    nodes, relationships = _projection_records()
    package = assemble_projection_package(
        kb_id=PROJ_KB,
        run=_fake_run(nodes, relationships),
        reference=_matching_reference(),
        evidence_bundle=_projection_evidence_bundle(),
        include_chunk_text=False,
    )

    with zipfile.ZipFile(io.BytesIO(package["content"])) as archive:
        chunk_rows = _read_jsonl(archive, "chunks.jsonl")
        manifest = json.loads(archive.read("manifest.json").decode("utf-8"))

    assert all(row["content"] is None for row in chunk_rows)
    assert all(row["content_sha256"] for row in chunk_rows)
    assert manifest["evidence"]["include_chunk_text"] is False
    assert any("本次未导出段落全文" in note for note in manifest["notes"])


class _FakeGraphRepository:
    """假 PG 仓储：异步生成器逐行产出，记录调用参数以证明导出是拉取式而非整表物化。"""

    def __init__(self, rows, chunks):
        self.rows = rows
        self.chunks = chunks
        self.evidence_calls: list[tuple] = []
        self.chunk_calls: list[tuple] = []

    async def iter_projection_evidence(self, kb_id, *, page_size=1000):
        self.evidence_calls.append((kb_id, page_size))
        for row in self.rows:
            yield row

    async def iter_projection_chunks(self, kb_id, chunk_ids, *, batch_size=500):
        requested = list(chunk_ids)
        self.chunk_calls.append((kb_id, requested, batch_size))
        for chunk in self.chunks:
            if chunk["chunk_id"] in set(requested):
                yield chunk


class _FakeReviewRepository:
    def __init__(self, decisions):
        self.decisions = decisions
        self.calls: list[str] = []

    async def list_decisions(self, kb_id):
        self.calls.append(kb_id)
        return self.decisions


async def test_export_projection_package_streams_evidence_from_repositories():
    nodes, relationships = _projection_records()
    transaction = _FakeTransaction(_fake_run(nodes, relationships))
    driver = _FakeDriver(transaction)
    bundle = _projection_evidence_bundle()
    graph_repository = _FakeGraphRepository(bundle["rows"], bundle["chunks"])
    review_repository = _FakeReviewRepository(bundle["decisions"])

    package = await export_projection_package(
        kb_id=PROJ_KB,
        driver=driver,
        reference=_matching_reference(),
        exported_by="tester",
        graph_repository=graph_repository,
        review_repository=review_repository,
        evidence_page_size=2,
        chunk_batch_size=2,
    )

    manifest = package["manifest"]
    with zipfile.ZipFile(io.BytesIO(package["content"])) as archive:
        assert set(archive.namelist()) == {
            "nodes.jsonl",
            "relationships.jsonl",
            "evidence.jsonl",
            "chunks.jsonl",
            "decisions.jsonl",
            "manifest.json",
        }
        evidence_rows = _read_jsonl(archive, "evidence.jsonl")
        assert (
            manifest["files"]["evidence.jsonl"]["sha256"] == hashlib.sha256(archive.read("evidence.jsonl")).hexdigest()
        )

    # 图成员仍在同一个已提交的显式读事务里读（查询顺序不变）
    assert transaction.queries == [
        PROJ_CYPHER["node_counts"],
        PROJ_CYPHER["relationship_counts"],
        PROJ_CYPHER["nodes"],
        PROJ_CYPHER["relationships"],
    ]
    assert transaction.committed and driver.session_instance.transactions_begun == 1
    assert graph_repository.evidence_calls == [(PROJ_KB, 2)]
    # chunk 只按证据引用到的 id 有序去重拉取（不是整库全表）
    assert graph_repository.chunk_calls == [(PROJ_KB, ["c1", "c2", "chunk_1"], 2)]
    assert review_repository.calls == [PROJ_KB]
    assert manifest["exported_by"] == "tester"
    assert manifest["evidence"]["coverage"]["evidence_complete"] is True
    assert len(evidence_rows) == 3


async def test_export_projection_package_without_evidence_skips_pg_reads():
    nodes, relationships = _projection_records()
    driver = _FakeDriver(_FakeTransaction(_fake_run(nodes, relationships)))
    graph_repository = _FakeGraphRepository([], [])
    review_repository = _FakeReviewRepository([])

    package = await export_projection_package(
        kb_id=PROJ_KB,
        driver=driver,
        reference=_matching_reference(),
        include_evidence=False,
        graph_repository=graph_repository,
        review_repository=review_repository,
    )

    with zipfile.ZipFile(io.BytesIO(package["content"])) as archive:
        assert set(archive.namelist()) == {"nodes.jsonl", "relationships.jsonl", "manifest.json"}

    assert package["manifest"]["schema_version"] == "neo4j-projection-jsonl-v1"
    assert graph_repository.evidence_calls == [] and review_repository.calls == []


async def test_export_projection_package_fails_closed_before_pg_read():
    nodes, relationships = _projection_records()
    run = _fake_run(nodes, relationships, node_counts={"total": 4, "entity": 3, "chunk": 1})
    driver = _FakeDriver(_FakeTransaction(run))
    graph_repository = _FakeGraphRepository([_mention_record()], [_chunk_record("c1")])

    with pytest.raises(RuntimeError, match="计数校验失败"):
        await export_projection_package(
            kb_id=PROJ_KB,
            driver=driver,
            reference=_matching_reference(),
            graph_repository=graph_repository,
            review_repository=_FakeReviewRepository([]),
        )

    assert graph_repository.evidence_calls == []


def test_evidence_workbook_appends_mention_quote_sheets(sample_snapshot):
    package = build_evidence_workbook(
        kb_id=KB,
        entities=sample_snapshot["entities"],
        triples=sample_snapshot["triples"],
        evidence=sample_snapshot["evidence"],
        aliases_by_entity=sample_snapshot["aliases"],
        mention_evidence=[
            _mention_record(target_id=T_PROMOTES),
            _mention_record(
                kind="entity",
                target_id=GENE_ENTITY_ID,
                chunk_id="chunk_1",
                edge_business_id=f"chunk_1->{GENE_ENTITY_ID}",
            ),
        ],
    )

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(package["content"]))
    assert workbook.sheetnames == ["实体", "三元组", "证据明细", "原文引文-三元组", "原文引文-实体"]
    triple_sheet = workbook["原文引文-三元组"]
    headers = [cell.value for cell in triple_sheet[1]]
    assert headers[:4] == ["triple_id", "source", "relation_type", "target"]
    for column in ("quote", "verification", "filename", "section", "review_status", "confidence", "trigger_term"):
        assert column in headers
    assert triple_sheet.max_row == 2
    assert triple_sheet.cell(row=2, column=1).value == T_PROMOTES
    assert triple_sheet.cell(row=2, column=2).value == "Os01g0100100(Gene)"
    triple_values = [cell.value for cell in triple_sheet[2]]
    assert CHUNK_QUOTE in triple_values and "OK" in triple_values
    entity_sheet = workbook["原文引文-实体"]
    assert entity_sheet.max_row == 2
    assert entity_sheet.cell(row=2, column=1).value == GENE_ENTITY_ID
    assert entity_sheet.cell(row=2, column=2).value == "Os01g0100100(Gene)"


def test_evidence_summary_header_is_ascii_and_reports_gaps():
    nodes, relationships = _projection_records()
    bundle = _projection_evidence_bundle()
    bundle["rows"].append(_mention_record(target_id="tri_8", chunk_id="c8", quote=None))
    package = assemble_projection_package(
        kb_id=PROJ_KB,
        run=_fake_run(nodes, relationships),
        reference=_matching_reference(),
        evidence_bundle=bundle,
    )

    header = evidence_summary_header(package)

    assert header == "evidence=4; ok=3; missing=1; chunks=3; decisions=1; complete=false"
    header.encode("latin-1")  # HTTP 头值必须可 latin-1 编码


async def test_write_jsonl_member_async_matches_sync_bytes():
    nodes, _ = _projection_records()
    sync_buffer = io.BytesIO()
    with zipfile.ZipFile(sync_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        sync_stat = write_jsonl_member(archive, "nodes.jsonl", iter(nodes), projection_node_row)

    async def records():
        for node in nodes:
            yield node

    async_buffer = io.BytesIO()
    with zipfile.ZipFile(async_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        # batch_rows=1 强制走多批落盘路径，验证分批不影响字节与统计
        async_stat = await write_jsonl_member_async(
            archive, "nodes.jsonl", records(), projection_node_row, batch_rows=1
        )

    assert async_stat == sync_stat
    with zipfile.ZipFile(io.BytesIO(async_buffer.getvalue())) as archive:
        assert archive.read("nodes.jsonl") == zipfile.ZipFile(io.BytesIO(sync_buffer.getvalue())).read("nodes.jsonl")
    evidence = _collect_evidence(
        [
            _mention_record(),  # RELATION 边 (tri_1, c1) 有引文
            _mention_record(chunk_id="c2", quote=None),  # 旧数据：边在、PG 无引文
            _mention_record(
                kind="entity", target_id="ent_gene", chunk_id="chunk_1", edge_business_id="chunk_1->ent_gene"
            ),  # MENTIONS 边有引文
            _mention_record(target_id="tri_9", review_status="REJECTED"),  # 边已删（预期，不算漂移）
            _mention_record(target_id="tri_8", review_status="CANDIDATE"),  # PG 有证据、投影无边（漂移）
        ]
    )

    coverage = evidence_coverage(_coverage_tally(), evidence)
    relation = coverage["relation_edges"]

    assert coverage["evidence_complete"] is False
    # 投影里有 RELATION 边 (tri_1, c2) 却找不到引文 → 覆盖缺口按 triple@chunk 报出
    assert relation["edges_without_evidence_count"] == 1
    assert relation["edges_without_evidence"] == ["tri_1@c2"]
    assert relation["evidence_without_edge"] == ["tri_8@c1", "tri_9@c1"]
    assert relation["evidence_without_edge_rejected"] == ["tri_9@c1"]
    assert relation["evidence_without_edge_orphan_count"] == 1
    assert relation["evidence_without_edge_orphan"] == ["tri_8@c1"]
    assert coverage["mention_edges"]["evidence_without_edge_count"] == 0
    assert evidence.verification == {"OK": 4, "MISSING": 1}


def test_evidence_coverage_is_complete_when_edges_and_quotes_align():
    evidence = _collect_evidence(
        [
            _mention_record(),
            _mention_record(chunk_id="c2"),
            _mention_record(
                kind="entity", target_id="ent_gene", chunk_id="chunk_1", edge_business_id="chunk_1->ent_gene"
            ),
        ]
    )

    coverage = evidence_coverage(_coverage_tally(), evidence)

    assert coverage["evidence_complete"] is True
    assert coverage["relation_edges"]["edges_without_evidence_count"] == 0
    assert coverage["relation_edges"]["evidence_without_edge_orphan_count"] == 0
    assert coverage["mention_edges"]["edges_without_evidence_count"] == 0
